#!/usr/bin/env python3
"""Build the verdict workbook of one persona's QA procedures: one row per step.

Reads docs/qa/<persona>/NNN-*.md, the same Markdown `make qa-docs` turns into
DOCX, and writes build/qa-docs/<persona>-qa-procedures.xlsx: a Read me sheet,
an Accounts sheet (the persona's README "Before you start", its accounts
table with each alias's address derived from the mailbox the tester types
once, and the fixture files), a Summary sheet that tallies the verdicts per
case by formula, and one sheet per procedure with a row per step (Section,
Case, Step, Step id, Action, Expected result, Pass/Fail, Notes). The
Pass/Fail cells take PASS, FAIL or N/A from a list; a step with no expected
result is setup and takes no verdict. The workbook stands on its own: a
tester needs nothing from the documents to know which account a step means.

    uv run --locked --group qa-docs python scripts/qa_sheets.py --persona governance
    uv run --locked --group qa-docs python scripts/qa_sheets.py --check --persona governance

`--check` parses the procedures, builds the workbook in memory and reads it
back, without writing a file; it is the guard that the Markdown still has the
shape this parser reads. See docs/how-to/publish-qa-procedures.md.
"""

from __future__ import annotations

import argparse
import io
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from publish_qa_docs import REPO_ROOT, REPO_URL, BuildInfo, _display, doc_name, source_files

STEP_HEADER = ["Step", "Action", "Expected result", "Pass/Fail", "Notes"]
COLUMNS = ["Section", "Case", "Step", "Step id", "Action", "Expected result", "Pass/Fail", "Notes"]
COLUMN_WIDTHS = [9, 8, 6, 9, 55, 70, 10, 30]
VERDICTS = ["PASS", "FAIL", "N/A"]
VERDICT_COLUMN = "G"
FIRST_DATA_ROW = 9
HEADER_ROW = FIRST_DATA_ROW - 1
DRIVE_FOLDER = "CarbonOS QA"


# --- the procedure as the Markdown states it -------------------------------------------


@dataclass
class Step:
    number: str
    action: str
    expected: str

    @property
    def is_setup(self) -> bool:
        return self.expected == ""


@dataclass
class Case:
    id: str
    title: str
    rationale: str = ""
    steps: list[Step] = field(default_factory=list)


@dataclass
class Section:
    letter: str
    title: str
    intro: str = ""
    cases: list[Case] = field(default_factory=list)


@dataclass
class Procedure:
    number: int
    title: str
    source: Path
    objective: str = ""
    covers: str = ""
    estimated: str = ""
    version: str = ""
    run_after: str = ""
    prerequisites: list[str] = field(default_factory=list)
    sections: list[Section] = field(default_factory=list)
    known_non_goals: str = ""
    change_notes: list[str] = field(default_factory=list)

    @property
    def cases(self) -> list[Case]:
        return [case for section in self.sections for case in section.cases]

    @property
    def steps(self) -> list[Step]:
        return [step for case in self.cases for step in case.steps]

    @property
    def verdict_steps(self) -> int:
        return sum(1 for step in self.steps if not step.is_setup)


# --- markdown ----------------------------------------------------------------------------

_BOLD = re.compile(r"\*\*(.+?)\*\*")
_CODE = re.compile(r"`([^`]*)`")
_LINK = re.compile(r"\[([^\]]+)\]\([^)]*\)")
_COMMENT = re.compile(r"<!--.*?-->", re.S)


def strip_markdown(text: str) -> str:
    """The cell's words without the Markdown that marked them: bold, code, links, escaped pipes."""
    text = _COMMENT.sub("", text)
    text = _LINK.sub(r"\1", text)
    text = _BOLD.sub(r"\1", text)
    text = _CODE.sub(r"\1", text)
    return text.replace("\\|", "|").strip()


def split_row(line: str) -> list[str]:
    """The cells of a pipe-table row, honouring `\\|` inside a cell."""
    body = line.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|") and not body.endswith("\\|"):
        body = body[:-1]
    return [cell.strip() for cell in re.split(r"(?<!\\)\|", body)]


def step_id(procedure: Procedure, case: Case, step: Step) -> str:
    return f"{procedure.number}.{case.id}.{step.number}"


class ParseError(SystemExit):
    def __init__(self, source: Path, line: int, message: str) -> None:
        super().__init__(f"{_display(source)}:{line}: {message}")


_H1 = re.compile(r"^# Procedure (\d+): (.+)$")
_SECTION = re.compile(r"^## ([A-Z])\. (.+)$")
_CASE = re.compile(r"^### ([A-Z]\d+[a-z]?)\. (.+)$")
_LEAD = re.compile(r"^\*\*(Objective|Covers|Estimated time|Procedure version|Run this procedure|Known non-goals)[.:]?\*\*\s*(.*)$")


def parse_procedure(source: Path) -> Procedure:  # noqa: C901 (one state machine, read top to bottom)
    """One generated procedure, read line by line in the order make qa-export writes it."""
    lines = source.read_text(encoding="utf-8").splitlines()
    procedure: Procedure | None = None
    section: Section | None = None
    case: Case | None = None
    mode = "head"  # head | prerequisites | body | signoff | changes
    table_header_seen = False
    paragraph: list[str] = []
    paragraph_target: str | None = None  # a procedure lead-in or "rationale" / "intro"
    case_ids: set[str] = set()

    def flush_paragraph() -> None:
        """Appends the collected lines to their target; the target stays until a heading or a table moves it."""
        nonlocal paragraph
        if paragraph and paragraph_target and procedure is not None:
            text = strip_markdown(" ".join(paragraph))
            if paragraph_target == "rationale" and case is not None:
                case.rationale = (case.rationale + " " + text).strip()
            elif paragraph_target == "intro" and section is not None:
                section.intro = (section.intro + " " + text).strip()
            elif paragraph_target not in ("rationale", "intro"):
                setattr(procedure, paragraph_target, (getattr(procedure, paragraph_target) + " " + text).strip())
        paragraph = []

    for index, raw in enumerate(lines, start=1):
        line = raw.rstrip()
        if _COMMENT.fullmatch(line.strip()):
            continue
        if h1 := _H1.match(line):
            procedure = Procedure(int(h1.group(1)), strip_markdown(h1.group(2)), source)
            continue
        if procedure is None:
            if line.strip():
                raise ParseError(source, index, "text before the '# Procedure N:' heading")
            continue

        if line == "## Prerequisites":
            flush_paragraph()
            mode, paragraph_target = "prerequisites", None
            continue
        if line == "## Sign-off":
            flush_paragraph()
            mode, case, section, paragraph_target = "signoff", None, None, None
            continue
        if line == "## Change notes":
            flush_paragraph()
            mode, paragraph_target = "changes", None
            continue
        if sec := _SECTION.match(line):
            flush_paragraph()
            if mode == "signoff" or mode == "changes":
                raise ParseError(source, index, "a section after the sign-off")
            mode = "body"
            section = Section(sec.group(1), strip_markdown(sec.group(2)))
            procedure.sections.append(section)
            case = None
            paragraph_target = "intro"
            continue
        if cs := _CASE.match(line):
            flush_paragraph()
            if section is None:
                raise ParseError(source, index, "a case outside a section")
            case_id = cs.group(1)
            if case_id in case_ids:
                raise ParseError(source, index, f"case {case_id} appears twice")
            case_ids.add(case_id)
            case = Case(case_id, strip_markdown(cs.group(2)))
            section.cases.append(case)
            table_header_seen = False
            paragraph_target = "rationale"
            continue

        if line.startswith("|"):
            cells = split_row(line)
            if mode == "signoff":
                continue  # the empty sign-off table: the workbook carries its own on Summary
            if mode != "body" or case is None:
                raise ParseError(source, index, "a table outside a case")
            if all(set(cell) <= {"-", ":", " "} and cell for cell in cells):
                continue  # the header separator
            if not table_header_seen:
                if cells != STEP_HEADER:
                    raise ParseError(source, index, f"step table header is {cells}, expected {STEP_HEADER}")
                table_header_seen = True
                flush_paragraph()
                paragraph_target = None
                continue
            if len(cells) != len(STEP_HEADER):
                raise ParseError(source, index, f"a step row with {len(cells)} cells")
            case.steps.append(Step(cells[0], strip_markdown(cells[1]), strip_markdown(cells[2])))
            continue

        if mode == "prerequisites":
            if line.startswith("- "):
                procedure.prerequisites.append(strip_markdown(line[2:]))
            elif line.startswith("  ") and procedure.prerequisites:
                procedure.prerequisites[-1] += " " + strip_markdown(line)
            continue
        if mode == "changes":
            if line.startswith("- "):
                procedure.change_notes.append(strip_markdown(line[2:]))
            elif line.strip() and procedure.change_notes:
                procedure.change_notes[-1] += " " + strip_markdown(line)
            continue

        if lead := _LEAD.match(line):
            flush_paragraph()
            paragraph_target = {
                "Objective": "objective",
                "Covers": "covers",
                "Estimated time": "estimated",
                "Procedure version": "version",
                "Run this procedure": "run_after",
                "Known non-goals": "known_non_goals",
            }[lead.group(1)]
            paragraph = [lead.group(2)]
            continue
        if not line.strip():
            flush_paragraph()
            continue
        if paragraph_target is not None:
            paragraph.append(line)
            continue
        raise ParseError(source, index, f"unexpected text: {line[:60]!r}")

    flush_paragraph()
    if procedure is None:
        raise ParseError(source, 1, "no '# Procedure N:' heading")
    if not procedure.sections:
        raise ParseError(source, len(lines), "no sections")
    for case in procedure.cases:
        if not case.steps:
            raise ParseError(source, len(lines), f"case {case.id} has no steps")
    return procedure


def parse_persona(persona: str) -> list[Procedure]:
    return [parse_procedure(path) for path in source_files(persona) if path.name != "README.md"]


# --- the persona's README: what a tester needs before the first step ---------------------


@dataclass
class Setup:
    """The parts of docs/qa/<persona>/README.md a tester needs beside the steps."""

    before: list[str] = field(default_factory=list)  # the "Before you start" bullets
    accounts_header: list[str] = field(default_factory=list)
    accounts: list[list[str]] = field(default_factory=list)  # the rows of "The accounts"
    accounts_note: str = ""  # the paragraph under the table ("Replace you+…@… with ...")
    fixtures_header: list[str] = field(default_factory=list)
    fixtures: list[list[str]] = field(default_factory=list)


_ALIAS = re.compile(r"you\+([a-z0-9]+)@")


def alias_of(account_cell: str) -> str | None:
    """'Ama Owusu, `you+ama@…`' -> 'ama'; None for an account with no alias (Admin A)."""
    hit = _ALIAS.search(account_cell)
    return hit.group(1) if hit else None


def parse_readme(source: Path) -> Setup:
    """The README's "Before you start" bullets and its accounts and fixture tables, as the workbook prints them."""
    setup = Setup()
    section = None
    paragraph: list[str] = []
    for index, raw in enumerate(source.read_text(encoding="utf-8").splitlines(), start=1):
        line = raw.rstrip()
        if line.startswith("## "):
            section = line[3:].strip()
            paragraph = []
            continue
        if section == "Before you start":
            if line.startswith("- "):
                setup.before.append(strip_markdown(line[2:]))
            elif line.startswith("  ") and setup.before:
                setup.before[-1] += " " + strip_markdown(line)
            continue
        if section in ("The accounts", "The fixture files"):
            header = setup.accounts_header if section == "The accounts" else setup.fixtures_header
            rows = setup.accounts if section == "The accounts" else setup.fixtures
            if line.startswith("|"):
                cells = split_row(line)
                if all(set(cell) <= {"-", ":", " "} and cell for cell in cells):
                    continue
                if not header:
                    header.extend(strip_markdown(cell) for cell in cells)
                elif len(cells) != len(header):
                    raise ParseError(source, index, f"a row with {len(cells)} cells under a {len(header)}-column table")
                else:
                    rows.append([strip_markdown(cell) for cell in cells])
            elif section == "The accounts" and setup.accounts:
                if line.strip():
                    paragraph.append(strip_markdown(line))
                elif paragraph:
                    setup.accounts_note = " ".join(paragraph)
                    paragraph = []
    if not setup.accounts:
        raise ParseError(source, 0, 'no "## The accounts" table')
    if not setup.before:
        raise ParseError(source, 0, 'no "## Before you start" bullets')
    return setup


def readme_path(persona: str) -> Path:
    return next(path for path in source_files(persona) if path.name == "README.md")


# --- the workbook --------------------------------------------------------------------------


def sheet_name(source: Path, taken: set[str]) -> str:
    """doc_name() made legal for a sheet: 31 characters, none of []:*?/\\, unique."""
    name = re.sub(r"[\[\]:*?/\\]", "", doc_name(source))[:31].rstrip()
    if name in taken:
        raise SystemExit(f"two procedures share the sheet name {name!r} after truncation")
    taken.add(name)
    return name


def quoted(name: str) -> str:
    return "'" + name.replace("'", "''") + "'"


MAILBOX_CELL = "B4"  # the Accounts sheet: the mailbox the tester reads, typed once


def build_workbook(procedures: list[Procedure], persona: str, info: BuildInfo, setup: Setup | None = None):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    bold = Font(bold=True)
    italic = Font(italic=True)
    title_font = Font(bold=True, size=14)
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")
    header_fill = PatternFill("solid", fgColor="D9E2EC")
    section_fill = PatternFill("solid", fgColor="BCCCDC")
    case_fill = PatternFill("solid", fgColor="E8EEF4")
    setup_fill = PatternFill("solid", fgColor="EDEDED")
    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    na_fill = PatternFill("solid", fgColor="D9D9D9")
    last_col = get_column_letter(len(COLUMNS))

    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    accounts = wb.create_sheet("Accounts")
    summary = wb.create_sheet("Summary")

    # one sheet per procedure; remember each case's verdict rows for the Summary formulas
    case_ranges: list[tuple[Procedure, Case, str, int, int]] = []  # procedure, case, sheet, first row, last row
    sheet_names: dict[int, str] = {}
    taken: set[str] = set()
    for procedure in procedures:
        name = sheet_name(procedure.source, taken)
        sheet_names[procedure.number] = name
        ws = wb.create_sheet(name)
        head = [
            (f"Procedure {procedure.number}: {procedure.title}", ""),
            ("Version", info.subtitle),
            ("Objective", procedure.objective),
            ("Covers", procedure.covers),
            ("Estimated time", procedure.estimated),
            ("Run this procedure", procedure.run_after),
            ("Prerequisites", "\n".join(f"- {p}" for p in procedure.prerequisites)),
        ]
        for row, (label, value) in enumerate(head, start=1):
            ws.cell(row=row, column=1, value=label).font = title_font if row == 1 else bold
            ws.cell(row=row, column=2, value=value).alignment = wrap
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=len(COLUMNS))
        ws.row_dimensions[7].height = max(15, 15 * len(procedure.prerequisites))
        for col, header in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=HEADER_ROW, column=col, value=header)
            cell.font, cell.fill = bold, header_fill
        for col, width in enumerate(COLUMN_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width

        validation = DataValidation(type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True)
        validation.error = "Type PASS, FAIL or N/A."
        validation.errorTitle = "Verdict"
        ws.add_data_validation(validation)

        row = FIRST_DATA_ROW
        for section in procedure.sections:
            cell = ws.cell(row=row, column=1, value=f"{section.letter}. {section.title}")
            cell.font, cell.fill = bold, section_fill
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLUMNS))
            row += 1
            if section.intro:
                ws.cell(row=row, column=1, value=section.intro).font = italic
                ws.cell(row=row, column=1).alignment = wrap
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLUMNS))
                row += 1
            for case in section.cases:
                cell = ws.cell(row=row, column=1, value=f"{case.id}. {case.title}")
                cell.font, cell.fill = bold, case_fill
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLUMNS))
                row += 1
                if case.rationale:
                    ws.cell(row=row, column=1, value=case.rationale).font = italic
                    ws.cell(row=row, column=1).alignment = wrap
                    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(COLUMNS))
                    row += 1
                first = row
                for step in case.steps:
                    values = [section.letter, case.id, step.number, step_id(procedure, case, step),
                              step.action, step.expected, None, None]
                    for col, value in enumerate(values, start=1):
                        cell = ws.cell(row=row, column=col, value=value)
                        cell.alignment = wrap if col in (5, 6, 8) else top
                        if step.is_setup:
                            cell.fill = setup_fill
                    if not step.is_setup:
                        validation.add(f"{VERDICT_COLUMN}{row}")
                    row += 1
                case_ranges.append((procedure, case, name, first, row - 1))
        last_step_row = row - 1

        row += 1
        ws.cell(row=row, column=1, value="Known non-goals").font = bold
        ws.cell(row=row, column=2, value=procedure.known_non_goals).alignment = wrap
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=len(COLUMNS))
        row += 1
        if procedure.change_notes:
            ws.cell(row=row, column=1, value="Change notes").font = bold
            for note in procedure.change_notes:
                ws.cell(row=row, column=2, value=note).alignment = wrap
                ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=len(COLUMNS))
                row += 1

        verdict_range = f"{VERDICT_COLUMN}{FIRST_DATA_ROW}:{VERDICT_COLUMN}{last_step_row}"
        ws.conditional_formatting.add(verdict_range, CellIsRule(operator="equal", formula=['"PASS"'], fill=pass_fill))
        ws.conditional_formatting.add(verdict_range, CellIsRule(operator="equal", formula=['"FAIL"'], fill=fail_fill))
        ws.conditional_formatting.add(verdict_range, CellIsRule(operator="equal", formula=['"N/A"'], fill=na_fill))
        ws.freeze_panes = f"A{FIRST_DATA_ROW}"
        ws.auto_filter.ref = f"A{HEADER_ROW}:{last_col}{last_step_row}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"

    # Summary: the sign-off per procedure, then the tally per case
    summary.cell(row=1, column=1, value=f"QA procedures: {persona}").font = title_font
    summary.cell(row=2, column=1, value=info.subtitle)
    row = 4
    for col, header in enumerate(["Procedure", "Procedure and version tested", "Tester", "Date", "Issues filed"], start=1):
        cell = summary.cell(row=row, column=col, value=header)
        cell.font, cell.fill = bold, header_fill
    for procedure in procedures:
        row += 1
        version = f"Procedure {procedure.number}, {info.subtitle}"
        if procedure.version:
            version += f"; procedure version {procedure.version.split('.')[0]}"
        summary.cell(row=row, column=1, value=f"{procedure.number:02d} {procedure.title}")
        summary.cell(row=row, column=2, value=version)
    row += 2
    tally_header = ["Procedure", "Case", "Title", "Verdict steps", "PASS", "FAIL", "N/A", "Blank", "Result"]
    for col, header in enumerate(tally_header, start=1):
        cell = summary.cell(row=row, column=col, value=header)
        cell.font, cell.fill = bold, header_fill
    first_tally = row + 1
    for procedure, case, name, first, last in case_ranges:
        row += 1
        ref = f"{quoted(name)}!${VERDICT_COLUMN}${first}:${VERDICT_COLUMN}${last}"
        verdicts = sum(1 for step in case.steps if not step.is_setup)
        summary.cell(row=row, column=1, value=procedure.number)
        summary.cell(row=row, column=2, value=case.id)
        summary.cell(row=row, column=3, value=case.title)
        summary.cell(row=row, column=4, value=verdicts)
        summary.cell(row=row, column=5, value=f'=COUNTIF({ref},"PASS")')
        summary.cell(row=row, column=6, value=f'=COUNTIF({ref},"FAIL")')
        summary.cell(row=row, column=7, value=f'=COUNTIF({ref},"N/A")')
        summary.cell(row=row, column=8, value=f"=D{row}-E{row}-F{row}-G{row}")
        summary.cell(row=row, column=9, value=f'=IF(F{row}>0,"FAIL",IF(H{row}>0,"",IF(E{row}>0,"PASS","N/A")))')
    last_tally = row
    row += 1
    summary.cell(row=row, column=3, value="Total").font = bold
    for col in range(4, 9):
        letter = get_column_letter(col)
        summary.cell(row=row, column=col, value=f"=SUM({letter}{first_tally}:{letter}{last_tally})").font = bold
    result_range = f"I{first_tally}:I{last_tally}"
    summary.conditional_formatting.add(result_range, CellIsRule(operator="equal", formula=['"PASS"'], fill=pass_fill))
    summary.conditional_formatting.add(result_range, CellIsRule(operator="equal", formula=['"FAIL"'], fill=fail_fill))
    for col, width in enumerate([32, 8, 60, 13, 8, 8, 8, 8, 10], start=1):
        summary.column_dimensions[get_column_letter(col)].width = width
    summary.freeze_panes = f"A{first_tally}"

    # Accounts: what the README says before the first step, with the alias addresses derived from one cell
    setup = setup or Setup()
    accounts.cell(row=1, column=1, value="Accounts, windows and files").font = title_font
    accounts.cell(row=2, column=1, value=info.subtitle)
    accounts.cell(row=4, column=1, value="The mailbox you read").font = bold
    mailbox = accounts[MAILBOX_CELL]
    mailbox.fill = PatternFill("solid", fgColor="FFF2CC")
    accounts.cell(row=4, column=3, value="Type it once, for example tester@gmail.com. Every alias address below "
                  "follows from it; plus-aliases of one Gmail address arrive in the base inbox.").alignment = wrap
    accounts.merge_cells(start_row=4, start_column=3, end_row=4, end_column=7)
    row = 6
    accounts.cell(row=row, column=1, value="Before you start").font = bold
    for bullet in setup.before:
        row += 1
        accounts.cell(row=row, column=1, value=f"- {bullet}").alignment = wrap
        accounts.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        accounts.row_dimensions[row].height = 15 * max(1, len(bullet) // 110 + 1)
    row += 2
    accounts.cell(row=row, column=1, value="The accounts").font = bold
    row += 1
    # the README's columns: Account first, a password column somewhere, the rest carried over as they are
    password_col = next((i for i, h in enumerate(setup.accounts_header) if h.lower().startswith("password")), None)
    rest = [i for i in range(1, len(setup.accounts_header)) if i != password_col]
    account_header = ["Sign in as", "Email", "Password", *(setup.accounts_header[i] for i in rest)]
    for col, header in enumerate(account_header, start=1):
        cell = accounts.cell(row=row, column=col, value=header)
        cell.font, cell.fill = bold, header_fill
    for account in setup.accounts:
        row += 1
        label = account[0]
        alias = alias_of(label)
        name = _ALIAS.split(label)[0].rstrip(", `").strip() if alias else label
        password = account[password_col] if password_col is not None else ""
        if alias:
            # =IF($B$4="","you+ama@…",SUBSTITUTE($B$4,"@","+ama@")): the placeholder until the mailbox is typed
            email = (f'=IF(${MAILBOX_CELL[0]}${MAILBOX_CELL[1:]}="","you+{alias}@…",'
                     f'SUBSTITUTE(${MAILBOX_CELL[0]}${MAILBOX_CELL[1:]},"@","+{alias}@"))')
        else:
            email = "the address you gave the engineering team"
        values = [name, email, password, *(account[i] for i in rest)]
        for col, value in enumerate(values, start=1):
            cell = accounts.cell(row=row, column=col, value=value)
            cell.alignment = wrap
    if setup.accounts_note:
        row += 1
        accounts.cell(row=row, column=1, value=setup.accounts_note).font = italic
        accounts.cell(row=row, column=1).alignment = wrap
        accounts.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    row += 1
    accounts.cell(row=row, column=1, value='Where a step says "as Ama" or "sign in as Admin", use that row\'s email and '
                  "password in the window the step names; the windows keep an administrator and a member apart.").alignment = wrap
    accounts.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    if setup.fixtures:
        row += 2
        accounts.cell(row=row, column=1, value="The fixture files").font = bold
        row += 1
        for col, header in enumerate(setup.fixtures_header, start=1):
            cell = accounts.cell(row=row, column=col, value=header)
            cell.font, cell.fill = bold, header_fill
        for fixture in setup.fixtures:
            row += 1
            for col, value in enumerate(fixture, start=1):
                accounts.cell(row=row, column=col, value=value).alignment = wrap
        row += 1
        accounts.cell(row=row, column=1, value=f"The files are in the fixtures folder beside this workbook in {DRIVE_FOLDER}. "
                      "Do not edit them: the procedures name their row numbers and totals.").alignment = wrap
        accounts.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    for col, width in enumerate([24, 34, 20, 34, 12, 60, 14], start=1):
        accounts.column_dimensions[get_column_letter(col)].width = width

    # Read me
    lines = [
        (f"QA procedures: {persona}", title_font),
        (info.subtitle, None),
        ("", None),
        ("How to use this workbook", bold),
        ("Start on the Accounts sheet: type the mailbox you read in its yellow cell, and every account's address, "
         "password and window is there. A step that says \"as Ama\" or \"sign in as Admin\" means that row.", None),
        ("Each procedure is a sheet, run them in order. Each row is one step: do the Action, check the Expected result, "
         "then pick PASS, FAIL or N/A in the Pass/Fail column.", None),
        ("A grey row is setup: it has no expected result and takes no verdict.", None),
        ("Write what you observed in Notes, above all for a FAIL: the message, the value, the screen.", None),
        ("The Summary sheet tallies the verdicts per case as you go, and carries the sign-off: fill in Tester, Date "
         "and Issues filed for each procedure when you finish it.", None),
        ("File one issue per failed case with the QA failure template: "
         f"{REPO_URL}/issues/new?template=qa-failure.yml", None),
        ("", None),
        ("Where things are", bold),
        (f"The procedures as documents, with the same steps: the {DRIVE_FOLDER} folder in Drive, in this version's subfolder, "
         "beside the fixture files the steps name.", None),
        (f"The source of this workbook: {REPO_URL}/tree/{info.ref}/docs/qa/{persona}", None),
        ("", None),
        ("Procedures in this workbook", bold),
    ]
    for procedure in procedures:
        lines.append((f"{sheet_names[procedure.number]}: {len(procedure.cases)} cases, {len(procedure.steps)} steps, "
                      f"{procedure.verdict_steps} with a verdict, about {procedure.estimated.rstrip('.')}", None))
    for row, (text, font) in enumerate(lines, start=1):
        cell = readme.cell(row=row, column=1, value=text)
        cell.alignment = wrap
        if font:
            cell.font = font
    readme.column_dimensions["A"].width = 120
    return wb


def counts_line(target: str, procedures: list[Procedure]) -> str:
    cases = sum(len(p.cases) for p in procedures)
    steps = sum(len(p.steps) for p in procedures)
    verdicts = sum(p.verdict_steps for p in procedures)
    return f"{target}  ->  {len(procedures)} procedures, {cases} cases, {steps} steps, {verdicts} verdict steps"


def build(out_dir: Path, persona: str, info: BuildInfo | None = None) -> Path:
    """Writes build/qa-docs/<persona>-qa-procedures.xlsx and returns its path."""
    info = info or BuildInfo.detect()
    procedures = parse_persona(persona)
    setup = parse_readme(readme_path(persona))
    target = out_dir.resolve() / f"{persona}-qa-procedures.xlsx"
    target.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(procedures, persona, info, setup).save(target)
    print(f"built {counts_line(_display(target), procedures)}")
    return target


def check(persona: str) -> None:
    """Parses, builds in memory and reads the workbook back; prints the counts, writes nothing."""
    from openpyxl import load_workbook

    procedures = parse_persona(persona)
    setup = parse_readme(readme_path(persona))
    buffer = io.BytesIO()
    build_workbook(procedures, persona, BuildInfo("check", "0000000", "1970-01-01"), setup).save(buffer)
    buffer.seek(0)
    wb = load_workbook(buffer)
    expected_sheets = len(procedures) + 3
    if len(wb.sheetnames) != expected_sheets:
        raise SystemExit(f"the workbook has {len(wb.sheetnames)} sheets, expected {expected_sheets}")
    # every alias account resolves its address from the mailbox cell; the seeded administrator has none to derive
    emails = [c.value for row in wb["Accounts"].iter_rows(min_col=2, max_col=2) for c in row
              if isinstance(c.value, str) and c.value.startswith("=IF(")]
    aliases = [a for a in (alias_of(account[0]) for account in setup.accounts) if a]
    if len(emails) != len(aliases):
        raise SystemExit(f"the Accounts sheet derives {len(emails)} addresses, expected one per alias ({len(aliases)})")
    summary = wb["Summary"]
    formulas = [c.value for row in summary.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=COUNTIF(")]
    if len(formulas) != 3 * sum(len(p.cases) for p in procedures):
        raise SystemExit(f"the Summary carries {len(formulas)} COUNTIF formulas, expected three per case")
    print(f"checked {counts_line(f'{persona} workbook', procedures)}")


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", type=Path, default=REPO_ROOT / "build" / "qa-docs", help="where the workbook goes")
    parser.add_argument("--persona", default="governance", help="the folder under docs/qa to export (default: governance)")
    parser.add_argument("--check", action="store_true", help="parse and build in memory only; write nothing")
    args = parser.parse_args(argv)
    if args.check:
        check(args.persona)
    else:
        build(args.out, args.persona)


if __name__ == "__main__":
    sys.exit(main())
